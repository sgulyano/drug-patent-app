"""Export the database to CSV (UTF-8 with BOM so Excel shows Thai correctly) and optionally one .xlsx workbook.

    python -m etl.export_csv --out exports/ [--xlsx] [--tables]
"""
import argparse
import csv
import io
from pathlib import Path

from sqlalchemy import text

from backend.db import ROOT, get_engine, init_db

VIEWS = {
    "patents": "v_patents",
    "drugs": "v_drugs",
    "drug_patent_matrix": "v_drug_patent_matrix",
    "assessments": "v_assessments",
}
TABLES = ["drug_product", "substance", "product_substance", "patent", "patent_link", "external_patent_ref",
          "assessment", "status_map"]


def query(sql: str, params: dict | None = None, db_path=None) -> tuple[list[str], list[tuple]]:
    with get_engine(db_path).connect() as conn:
        res = conn.execute(text(sql), params or {})
        return list(res.keys()), [tuple(r) for r in res.fetchall()]


def to_csv_bytes(cols: list[str], rows: list[tuple]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(cols)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def view_sql(name: str) -> str:
    if name in VIEWS:
        return f"SELECT * FROM {VIEWS[name]}"
    if name in TABLES:
        return f"SELECT * FROM {name}"
    raise KeyError(name)


def export(out: Path, xlsx: bool = False, tables: bool = False, db_path=None) -> list[Path]:
    init_db(db_path)
    out.mkdir(parents=True, exist_ok=True)
    names = list(VIEWS) + (TABLES if tables else [])
    written, sheets = [], {}
    for name in names:
        cols, rows = query(view_sql(name), db_path=db_path)
        path = out / f"{name}.csv"
        path.write_bytes(to_csv_bytes(cols, rows))
        written.append(path)
        sheets[name] = (cols, rows)
    if xlsx:
        written.append(write_xlsx(out / "patent_landscape.xlsx", sheets))
    return written


def write_xlsx(path: Path, sheets: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, (cols, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(cols)
        for r in rows:
            ws.append([v if not isinstance(v, str) else ILLEGAL_CHARACTERS_RE.sub("", v)[:32000] for v in r])
        for c in ws[1]:
            c.font = Font(bold=True)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(cols, 1):
            width = max([len(str(col))] + [min(len(str(r[i - 1] or "")), 60) for r in rows[:200]])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(width + 2, 10), 60)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=False, vertical="top")
    wb.save(path)
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "exports"))
    ap.add_argument("--xlsx", action="store_true")
    ap.add_argument("--tables", action="store_true", help="also dump every raw table")
    a = ap.parse_args()
    for p in export(Path(a.out), a.xlsx, a.tables):
        print(p)
